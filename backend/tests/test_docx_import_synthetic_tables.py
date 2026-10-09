"""Synthetic multi-type Word document parsing regression for DocxImportService.

Ported from the retired perf fixture pipeline (test_perf_fixture.py /
generate_perf_fixture.py at f68ebe1): builds a synthetic .docx with two
preface tables plus form tables covering every field type, then asserts what
`DocxImportService.parse_full` extracts — preface-table skipping, form/field
counts, parsed field-type proportions, choice options, log rows, label
defaults, and numeric unit extraction.
"""

from __future__ import annotations

from collections import Counter
from pathlib import Path

from docx import Document

from src.services.docx_import_service import DocxImportService

SYN_SKIPPED_TABLE_COUNT = 2
SYN_FORM_TABLE_COUNT = 4
SYN_FIELDS_PER_FORM = 40

SYN_FIELD_TYPE_SEQUENCE = (
    ["文本"] * 8
    + ["数值"] * 5
    + ["日期"] * 4
    + ["日期时间"] * 3
    + ["时间"] * 2
    + ["单选"] * 4
    + ["多选"] * 4
    + ["单选（纵向）"] * 3
    + ["多选（纵向）"] * 3
    + ["标签"] * 2
    + ["日志行"] * 2
)
SYN_FIELD_TYPE_COUNTS = dict(Counter(SYN_FIELD_TYPE_SEQUENCE))
# parse_full emits log rows as {"type": "log_row"}; every other row keeps field_type.
SYN_PARSED_FIELD_TYPE_COUNTS = {
    "log_row" if field_type == "日志行" else field_type: count for field_type, count in SYN_FIELD_TYPE_COUNTS.items()
}
CHOICE_FIELD_TYPES = ("单选", "多选", "单选（纵向）", "多选（纵向）")


def _syn_docx_label_for(field_type: str, form_index: int, field_index: int) -> str:
    if field_type == "日志行":
        return f"SYN_说明行_{form_index:02d}_{field_index:02d}"
    return f"SYN_{field_type}_字段_{form_index:02d}_{field_index:02d}"


# Per-field-type value templates; `{suffix}` is substituted where present,
# constant templates pass through `.format` unchanged.
SYN_VALUE_TEMPLATES = {
    "文本": "",
    "数值": "|__|__|.|__|SYN单位",
    "日期": "|__|__|__|__|年|__|__|月|__|__|日",
    "日期时间": "|__|__|__|__|年|__|__|月|__|__|日 |__|__|:|__|__| SYN_DT_ASCII",
    "时间": "|__|__|:|__|__|",
    "单选": "○SYN_选项A_{suffix}  ○SYN_选项B_{suffix}",
    "多选": "□SYN_选项A_{suffix}  □SYN_选项B_{suffix}",
    "单选（纵向）": "○SYN_纵向选项A_{suffix}\n○SYN_纵向选项B_{suffix}",
    "多选（纵向）": "□SYN_纵向选项A_{suffix}\n□SYN_纵向选项B_{suffix}",
    "标签": "SYN_合成标签正文_{suffix}_ASCII_CJK",
    "日志行": "以下为log行",
}


def _syn_docx_value_for(field_type: str, form_index: int, field_index: int) -> str:
    suffix = f"{form_index:02d}_{field_index:02d}"
    template = SYN_VALUE_TEMPLATES.get(field_type, "SYN_PLACEHOLDER_{suffix}")
    return template.format(suffix=suffix)


def _add_syn_preface_table(document: Document, title: str) -> None:
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = title
    table.cell(0, 1).text = "SYN_SYNTHETIC"


def _build_synthetic_docx(docx_path: Path) -> dict[str, int]:
    document = Document()
    _add_syn_preface_table(document, "SYN_封面信息表")
    _add_syn_preface_table(document, "SYN_访视分布图")

    for form_index in range(1, SYN_FORM_TABLE_COUNT + 1):
        document.add_paragraph(f"SYN_表单_{form_index:02d}", style="List Paragraph")
        table = document.add_table(rows=0, cols=2)
        for field_index, field_type in enumerate(SYN_FIELD_TYPE_SEQUENCE, start=1):
            row = table.add_row()
            row.cells[0].text = _syn_docx_label_for(field_type, form_index, field_index)
            row.cells[1].text = _syn_docx_value_for(field_type, form_index, field_index)

    document.save(docx_path)
    return {
        "physical_tables": len(document.tables),
        "form_tables": SYN_FORM_TABLE_COUNT,
    }


def test_parse_full_skips_preface_tables_and_pairs_every_form_title_with_its_table(tmp_path: Path) -> None:
    # Arrange
    docx_path = tmp_path / "synthetic.docx"
    built = _build_synthetic_docx(docx_path)

    # Act
    parsed_forms = DocxImportService.parse_full(str(docx_path))

    # Assert
    assert built["physical_tables"] == SYN_SKIPPED_TABLE_COUNT + SYN_FORM_TABLE_COUNT
    assert len(parsed_forms) == built["form_tables"]
    assert [form["name"] for form in parsed_forms] == [
        f"SYN_表单_{index:02d}" for index in range(1, SYN_FORM_TABLE_COUNT + 1)
    ]
    assert all(len(form["fields"]) == SYN_FIELDS_PER_FORM for form in parsed_forms)


def test_parse_full_reads_every_field_type_in_expected_proportions(tmp_path: Path) -> None:
    # Arrange
    docx_path = tmp_path / "synthetic.docx"
    _build_synthetic_docx(docx_path)

    # Act
    parsed_forms = DocxImportService.parse_full(str(docx_path))

    # Assert
    assert sum(SYN_PARSED_FIELD_TYPE_COUNTS.values()) == SYN_FIELDS_PER_FORM
    for form in parsed_forms:
        counts = Counter(field.get("field_type", field.get("type", "未知")) for field in form["fields"])
        assert counts == SYN_PARSED_FIELD_TYPE_COUNTS


def test_parse_full_parses_two_options_for_every_choice_field(tmp_path: Path) -> None:
    # Arrange
    docx_path = tmp_path / "synthetic.docx"
    _build_synthetic_docx(docx_path)

    # Act
    parsed_forms = DocxImportService.parse_full(str(docx_path))

    # Assert
    choice_fields_per_form = sum(SYN_FIELD_TYPE_COUNTS[choice] for choice in CHOICE_FIELD_TYPES)
    for form in parsed_forms:
        choice_fields = [field for field in form["fields"] if field.get("field_type") in CHOICE_FIELD_TYPES]
        assert len(choice_fields) == choice_fields_per_form
        for field in choice_fields:
            assert len(field["options"]) == 2
            assert all(option["decode"].startswith("SYN_") for option in field["options"])


def test_parse_full_flags_log_rows_and_keeps_label_defaults_and_numeric_units(tmp_path: Path) -> None:
    # Arrange
    docx_path = tmp_path / "synthetic.docx"
    _build_synthetic_docx(docx_path)

    # Act
    parsed_forms = DocxImportService.parse_full(str(docx_path))

    # Assert
    for form in parsed_forms:
        fields = form["fields"]
        log_rows = [field for field in fields if field.get("type") == "log_row"]
        assert len(log_rows) == SYN_FIELD_TYPE_COUNTS["日志行"]
        label_fields = [field for field in fields if field.get("field_type") == "标签"]
        assert len(label_fields) == SYN_FIELD_TYPE_COUNTS["标签"]
        assert all(field["default_value"].startswith("SYN_合成标签正文_") for field in label_fields)
        numeric_fields = [field for field in fields if field.get("field_type") == "数值"]
        assert len(numeric_fields) == SYN_FIELD_TYPE_COUNTS["数值"]
        assert all(field["unit_symbol"] == "SYN单位" for field in numeric_fields)
